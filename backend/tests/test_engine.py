"""The engine on unrelated schemas - nothing about columns is hard-coded."""
import datetime as dt
import os
import random
import shutil
import tempfile
import time
from pathlib import Path

os.environ.setdefault("DN_STORAGE_DIR", tempfile.mkdtemp(prefix="dn-test-"))

import pandas as pd
import pytest
from fastapi.testclient import TestClient

from datanarrator import catalog, config, service
from datanarrator.connectors import sources

SAMPLES = Path(__file__).resolve().parents[2] / "samples"


def _import(path: Path, owner: str) -> list[dict]:
    dest = config.UPLOADS_DIR / path.name
    shutil.copy(path, dest)
    ph = catalog.create(owner, path.name, "import", "upload", status="importing")
    service._run_import(owner, ph["id"], path.name, "upload",
                        lambda con: sources.local_relations(con, dest, path.name), [dest])
    return catalog.list_for(owner)


def _roles(ds) -> dict:
    return {c["name"]: c["role"] for c in catalog.load_json(ds["id"], "profile")["columns"]}


@pytest.fixture(scope="module", autouse=True)
def _init():
    catalog.init()


def test_hr_dataset(tmp_path):
    rnd = random.Random(1)
    depts = ["Engineering", "Sales", "HR", "Finance", "Support"]
    df = pd.DataFrame({
        "Employee Code": [f"EMP{i:05d}" for i in range(800)],
        "Full Name": [f"Person {i}" for i in range(800)],
        "Department": [rnd.choice(depts) for _ in range(800)],
        "Joining Date": [(dt.date(2019, 1, 1) + dt.timedelta(days=rnd.randint(0, 1800))).strftime("%d/%m/%Y")
                         for _ in range(800)],
        "Monthly Salary": [rnd.randint(30, 200) * 1000 for _ in range(800)],
        "Performance Rating": [rnd.randint(1, 5) for _ in range(800)],
        "Attrition": [rnd.choice(["Yes", "No", "No", "No"]) for _ in range(800)],
    })
    df.loc[df["Department"] == "Engineering", "Monthly Salary"] *= 2
    path = tmp_path / "employees.csv"
    df.to_csv(path, index=False)
    ds = _import(path, "hr")[0]
    roles = _roles(ds)
    assert roles["Employee Code"] == "id"
    assert roles["Department"] == "dimension" and roles["Attrition"] == "dimension"
    assert roles["Joining Date"] == "time", "dd/mm/yyyy text parsed as dates"
    assert roles["Monthly Salary"] == "measure"
    res = catalog.load_json(ds["id"], "insights")
    assert any(i["title"] == "Monthly Salary by Department" and "Engineering" in i["summary"]
               for i in res["insights"])


def test_iot_dataset(tmp_path):
    rnd = random.Random(2)
    start = dt.datetime(2025, 3, 1)
    rows = []
    for h in range(24 * 60):
        for dev in ("sensor-a", "sensor-b", "sensor-c"):
            temp = 22 + 0.01 * h + rnd.random()
            rows.append({"device": dev, "ts": start + dt.timedelta(hours=h), "temperature_c": temp,
                         "power_kwh": temp * 1.5 + rnd.random()})
    rows[100]["temperature_c"] = 400.0  # a faulty reading
    path = tmp_path / "telemetry.parquet"
    pd.DataFrame(rows).to_parquet(path)
    ds = _import(path, "iot")[0]
    roles = _roles(ds)
    assert roles["ts"] == "time" and roles["device"] == "dimension"
    res = catalog.load_json(ds["id"], "insights")
    types = {i["type"] for i in res["insights"]}
    assert {"trend", "correlation", "outlier"} <= types


def test_messy_excel_end_to_end(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Q3 Sales"
    ws.merge_cells("A1:D1")
    ws.append(["Zone", "Item", "Revenue", "Margin"])
    for z, i, r, m in [("East", "Pens", "₹1,200", "(150)"), (None, "Ink", "₹2,500", "300"),
                       ("West", "Pens", "₹840", "90"), (None, "Ink", "₹4,500", "900"),
                       ("Total", None, "₹9,040", "1,140")]:
        ws.append([z, i, r, m])
    path = tmp_path / "q3.xlsx"
    wb.save(path)
    ds = _import(path, "xl")[0]
    prof = catalog.load_json(ds["id"], "profile")
    types = {c["name"]: c["type"] for c in prof["columns"]}
    assert types["Revenue"] in ("BIGINT", "DOUBLE") and types["Margin"] in ("BIGINT", "DOUBLE")
    assert prof["row_count"] == 4
    res = catalog.load_json(ds["id"], "insights")
    rev = next(k for k in res["kpis"] if k["label"] == "Revenue")
    assert rev["value"] == 9040, "₹ text summed correctly and the Total row excluded"


def test_sample_tables_are_joined():
    for f in sorted(SAMPLES.glob("*.csv")):
        _import(f, "sample")
    service.refresh_models("sample")
    models = catalog.list_for("sample", "model")
    assert [m["name"] for m in models] == ["Order Details + List of Orders"]
    res = catalog.load_json(models[0]["id"], "insights")
    assert any(i["type"] == "trend" for i in res["insights"]), "dates came from the joined table"
    assert any(a["priority"] == "high" for a in res["actions"])
    orders = next(d for d in catalog.list_for("sample", "table") if d["name"] == "List of Orders")
    assert orders["row_count"] == 500, "blank CSV rows dropped"


def test_api_upload_flow():
    from datanarrator.api.main import app
    with TestClient(app) as client:
        with open(SAMPLES / "Order Details.csv", "rb") as fh:
            r = client.post("/api/sources/upload", files={"files": ("Order Details.csv", fh, "text/csv")})
        assert r.status_code == 202
        for _ in range(100):
            ds = [d for d in client.get("/api/datasets").json()["datasets"] if d["name"] == "Order Details"]
            if ds and ds[0]["status"] == "ready":
                break
            time.sleep(0.1)
        assert ds and ds[0]["status"] == "ready"
        body = client.get(f"/api/datasets/{ds[0]['id']}/insights").json()
        assert body["kpis"] and body["insights"] and body["actions"]
        assert len(client.get(f"/api/datasets/{ds[0]['id']}/preview?limit=5").json()["rows"]) == 5
