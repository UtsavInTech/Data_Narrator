"""MIS flow: login by role, entry drafts, commit with Employee ID + password, living tables,
templates, daily report, write-back."""
import datetime as dt
import io
import os
import subprocess
import tempfile
import time

os.environ.setdefault("DN_STORAGE_DIR", tempfile.mkdtemp(prefix="dn-test-"))

import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from datanarrator.api.main import app


def _wait_ready(client, name, rows=None):
    for _ in range(150):
        ds = [d for d in client.get("/api/datasets").json()["datasets"] if d["name"] == name]
        if ds and ds[0]["status"] == "ready" and (rows is None or ds[0]["row_count"] == rows):
            return ds[0]
        time.sleep(0.1)
    raise AssertionError(f"{name} not ready: {ds}")


def _sales_csv(days: list[dt.date]) -> bytes:
    rows = [{"Date": d.isoformat(), "Branch": b, "Units": 10 + i, "Revenue": 1000.0 + 100 * i}
            for d in days for i, b in enumerate(["Delhi", "Pune", "Kochi"])]
    return pd.DataFrame(rows).to_csv(index=False).encode()


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        from datanarrator import auth
        if auth.setup_needed():
            c.post("/api/auth/setup", json={"employee_id": "ADM1", "name": "Admin", "password": "secret123"})
        token = c.post("/api/auth/login", json={"employee_id": "ADM1", "password": "secret123"}).json()["token"]
        c.headers["Authorization"] = f"Bearer {token}"
        c.post("/api/users", json={"employee_id": "OP7", "name": "Ravi", "password": "ravi1234", "role": "operator"})
        c.post("/api/users", json={"employee_id": "MG2", "name": "Meera", "password": "meera123", "role": "manager"})
        today = dt.date.today()
        c.post("/api/sources/upload", files={"files": ("daily_sales.csv", _sales_csv(
            [today - dt.timedelta(days=k) for k in range(3, 0, -1)]), "text/csv")})
        _wait_ready(c, "Daily sales")
        yield c


def _as(client, emp, pw):
    t = client.post("/api/auth/login", json={"employee_id": emp, "password": pw}).json()["token"]
    return {"Authorization": f"Bearer {t}"}


def test_login_required_and_roles(client):
    assert client.get("/api/datasets", headers={"Authorization": "Bearer nope"}).status_code == 401
    op = _as(client, "OP7", "ravi1234")
    assert client.get("/api/users", headers=op).status_code == 403
    assert client.get("/api/mis/reports", headers=op).status_code == 403
    mg = _as(client, "MG2", "meera123")
    assert client.get("/api/mis/reports", headers=mg).status_code == 200
    assert client.post("/api/mis/x/entries", json={"data": {}}, headers=mg).status_code == 403


def test_form_is_learnt_from_data(client):
    ds = _wait_ready(client, "Daily sales")
    fields = {f["name"]: f for f in client.get(f"/api/mis/{ds['id']}/form").json()["fields"]}
    assert fields["Branch"]["kind"] == "select" and fields["Branch"]["options"] == ["Delhi", "Kochi", "Pune"]
    assert fields["Date"]["kind"] == "date" and fields["Units"]["min"] == 0 and fields["Units"]["required"]


def test_entry_draft_commit_flow(client):
    ds = _wait_ready(client, "Daily sales")
    op = _as(client, "OP7", "ravi1234")
    base = f"/api/mis/{ds['id']}"
    today = dt.date.today().isoformat()
    bad = client.post(f"{base}/entries", json={"data": {"Date": today, "Branch": "Nashik", "Units": -4, "Revenue": "1,200"}},
                      headers=op).json()
    assert bad["errors"] == {"Units": "can't be negative"}
    assert "new" in bad["warnings"]["Branch"]
    r = client.post(f"{base}/commit", json={"employee_id": "OP7", "password": "ravi1234"}, headers=op)
    assert r.status_code == 422, "errors block the commit"
    client.put(f"{base}/entries/{bad['id']}", json={"data": {"Date": today, "Branch": "Nashik", "Units": 4, "Revenue": "1,200"}},
               headers=op)
    client.post(f"{base}/entries", json={"data": {"Date": today, "Branch": "Delhi", "Units": 9, "Revenue": 900}}, headers=op)
    assert client.post(f"{base}/commit", json={"employee_id": "OP7", "password": "wrong"}, headers=op).status_code == 401
    res = client.post(f"{base}/commit", json={"employee_id": "OP7", "password": "ravi1234"}, headers=op).json()
    assert res["rows"] == 2 and res["committed_by"]["employee_id"] == "OP7"
    assert client.get(f"{base}/entries", headers=op).json()["entries"] == []
    ds = _wait_ready(client, "Daily sales", rows=11)
    batches = client.get(f"/api/datasets/{ds['id']}").json()["batches"]
    assert batches[0]["kind"] == "entry" and batches[0]["created_by"] == "OP7"


def test_same_layout_upload_appends_and_duplicate_is_refused(client):
    today = dt.date.today()
    data = _sales_csv([today])
    r = client.post("/api/sources/upload", files={"files": ("daily_sales_today.csv", data, "text/csv")})
    assert r.status_code == 202
    ds = _wait_ready(client, "Daily sales", rows=14)
    assert not [d for d in client.get("/api/datasets").json()["datasets"] if d["name"] == "daily_sales_today"]
    again = client.post("/api/sources/upload", files={"files": ("copy.csv", data, "text/csv")})
    assert again.status_code == 409 and "already added" in again.json()["detail"]


def test_template_has_dropdowns(client):
    ds = _wait_ready(client, "Daily sales")
    r = client.get(f"/api/mis/{ds['id']}/template.xlsx")
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.worksheets[0]
    assert [c.value for c in ws[1]] == ["Date", "Branch", "Units", "Revenue"]
    assert any("Lists" in (dv.formula1 or "") for dv in ws.data_validations.dataValidation)


def test_daily_report(client):
    ds = _wait_ready(client, "Daily sales")
    mg = _as(client, "MG2", "meera123")
    r = client.get(f"/api/mis/{ds['id']}/daily-report.xlsx", headers=mg)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "By Branch", "Rows", "Checks"]
    summary = {row[0]: row for row in wb["Summary"].iter_rows(min_row=5, values_only=True) if row[0]}
    assert summary["Records"][1] == 5, "3 uploaded + 2 entered today"
    checks = [row[1] for row in wb["Checks"].iter_rows(min_row=2, values_only=True)]
    assert any("Nashik" in (c or "") for c in checks), "a new branch is flagged"
    assert client.get("/api/mis/reports", headers=mg).json()["reports"]


def _pg_available() -> bool:
    try:
        return "dn-pg" in subprocess.run(["docker", "ps", "--format", "{{.Names}}"], capture_output=True,
                                          text=True, timeout=5).stdout
    except Exception:
        return False


@pytest.mark.skipif(not _pg_available(), reason="needs the dn-pg test container")
def test_writeback_to_postgres(client):
    subprocess.run(["docker", "exec", "dn-pg", "psql", "-q", "-U", "postgres", "-d", "shop", "-c",
                    "DROP TABLE IF EXISTS grn; CREATE TABLE grn (grn_no int PRIMARY KEY, received date, sku text, qty int);"
                    "INSERT INTO grn SELECT i, current_date - i, 'SKU-' || (i % 5), 10 + i FROM generate_series(1, 60) i;"],
                   check=True)
    client.post("/api/sources/database", json={"type": "postgres", "host": "localhost", "port": 55432, "database": "shop",
                                               "user": "postgres", "password": "dnpass", "tables": ["grn"],
                                               "allow_writeback": True})
    ds = _wait_ready(client, "grn")
    assert ds["writeback"]
    base = f"/api/mis/{ds['id']}"
    dup = client.post(f"{base}/entries", json={"data": {"grn_no": 5, "received": dt.date.today().isoformat(),
                                                        "sku": "SKU-1", "qty": 3}}).json()
    assert dup["errors"].get("grn_no") == "5 already exists"
    client.put(f"{base}/entries/{dup['id']}", json={"data": {"grn_no": 9001, "received": dt.date.today().isoformat(),
                                                             "sku": "SKU-1", "qty": 3}})
    res = client.post(f"{base}/commit", json={"employee_id": "ADM1", "password": "secret123", "push": True}).json()
    assert res["pushed"] is True
    out = subprocess.run(["docker", "exec", "dn-pg", "psql", "-tA", "-U", "postgres", "-d", "shop", "-c",
                          "SELECT qty FROM grn WHERE grn_no = 9001"], capture_output=True, text=True).stdout.strip()
    assert out == "3"
