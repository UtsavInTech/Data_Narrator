"""Read real-world Excel workbooks the way an analyst would.

Handles: several sheets, several tables per sheet, title rows above the
header, merged cells, two-row headers, "Grand Total"/subtotal rows, repeated
header rows, report-style grouping (value only on the first row of a group),
month/year columns laid out sideways (unpivoted to rows), Excel error cells,
hidden sheets, defined Excel Tables, and identical sheets per month/region
(stacked into one table). Every decision is recorded in a per-table report.
"""
import datetime as dt
import re
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries

EXCEL_EXTS = {".xlsx", ".xlsm", ".xltx", ".xltm", ".xls", ".xlsb", ".ods"}
BIG_FILE_BYTES = 40 * 1024 * 1024

_ERROR_VALUES = {"#N/A", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}
_NULL_TEXT = {"", "-", "--", "—", "n/a", "na", "nan", "null", "none", "nil"}
_TOTAL_RE = re.compile(r"^\s*(grand\s+total|sub\s*-?\s*total|total|totals|sum|overall|net\s+total)\b", re.I)
_MONTHS = ("jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec")
_PERIOD_RE = re.compile(
    r"^((%s)[a-z]*[\s\-'/,]*(\d{2}|\d{4})?|(19|20)\d{2}|q[1-4][\s\-']*((19|20)?\d{2})?|"
    r"(fy|cy)[\s\-']?\d{2,4}(-\d{2,4})?|h[12][\s\-']*\d{2,4}|week\s*\d+|wk\s*\d+)$" % "|".join(_MONTHS),
    re.I,
)


@dataclass
class ParsedTable:
    name: str
    frame: pd.DataFrame
    report: dict = field(default_factory=dict)


# ---------------------------------------------------------------- cell helpers

def _empty(v) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "") or (isinstance(v, float) and v != v)


def _is_text(v) -> bool:
    return isinstance(v, str) and not _empty(v)


def _is_number(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool) and v == v


def _clean_cell(v, stats: dict):
    if isinstance(v, str):
        s = v.strip()
        if s.upper() in _ERROR_VALUES:
            stats["error_cells"] = stats.get("error_cells", 0) + 1
            return None
        if s.lower() in _NULL_TEXT:
            return None
        return s
    if isinstance(v, float) and v != v:
        return None
    return v


def _header_text(v) -> str:
    if isinstance(v, dt.datetime):
        return v.strftime("%b %Y") if v.day == 1 else v.strftime("%Y-%m-%d")
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip()


def _looks_period(v) -> bool:
    if isinstance(v, (dt.date, dt.datetime)):
        return True
    if _is_number(v):
        return 1900 <= v <= 2100 and float(v).is_integer()
    return _is_text(v) and bool(_PERIOD_RE.match(v.strip()))


# ---------------------------------------------------------------- loading

@dataclass
class _Sheet:
    name: str
    grid: list[list]
    hidden: bool = False
    tables: list[tuple[str, tuple[int, int, int, int]]] = field(default_factory=list)  # name, (r0,c0,r1,c1) 0-based incl.


def _trim(grid: list[list]) -> list[list]:
    while grid and all(_empty(v) for v in grid[-1]):
        grid.pop()
    width = 0
    for row in grid:
        for i in range(len(row) - 1, -1, -1):
            if not _empty(row[i]):
                width = max(width, i + 1)
                break
    return [list(row[:width]) + [None] * (width - len(row[:width])) for row in grid]


def _load_openpyxl(path: Path) -> list[_Sheet]:
    big = path.stat().st_size > BIG_FILE_BYTES
    wb = load_workbook(path, data_only=True, read_only=big)
    sheets = []
    for ws in wb.worksheets:
        grid = _trim([list(r) for r in ws.iter_rows(values_only=True)])
        if not big:
            # Merged cells: copy the top-left value into the whole range.
            for rng in ws.merged_cells.ranges:
                r0, c0 = rng.min_row - 1, rng.min_col - 1
                if r0 >= len(grid) or c0 >= len(grid[r0]):
                    continue
                v = grid[r0][c0]
                for r in range(r0, min(rng.max_row, len(grid))):
                    for c in range(c0, min(rng.max_col, len(grid[r]))):
                        grid[r][c] = v
        tables = []
        if not big:
            for tname, tbl in ws.tables.items():
                ref = tbl.ref if hasattr(tbl, "ref") else tbl
                c0, r0, c1, r1 = range_boundaries(ref)
                tables.append((tname, (r0 - 1, c0 - 1, r1 - 1, c1 - 1)))
        sheets.append(_Sheet(ws.title, grid, getattr(ws, "sheet_state", "visible") != "visible", tables))
    wb.close()
    return sheets


def _load_pandas(path: Path) -> list[_Sheet]:
    engine = {".xls": "xlrd", ".xlsb": "pyxlsb", ".ods": "odf"}.get(path.suffix.lower())
    books = pd.read_excel(path, sheet_name=None, header=None, engine=engine)
    return [_Sheet(str(name), _trim(df.astype(object).where(df.notna(), None).values.tolist()))
            for name, df in books.items()]


# ---------------------------------------------------------------- block detection

def _blocks(grid: list[list]) -> list[tuple[int, int, int, int]]:
    """Split a sheet into rectangular regions separated by empty rows/columns."""
    def split_rows(r0, r1, c0, c1):
        out, start = [], None
        for r in range(r0, r1 + 1):
            filled = any(not _empty(grid[r][c]) for c in range(c0, c1 + 1))
            if filled and start is None:
                start = r
            if not filled and start is not None:
                out.append((start, r - 1))
                start = None
        if start is not None:
            out.append((start, r1))
        return out

    def split_cols(r0, r1, c0, c1):
        out, start = [], None
        for c in range(c0, c1 + 1):
            filled = any(not _empty(grid[r][c]) for r in range(r0, r1 + 1))
            if filled and start is None:
                start = c
            if not filled and start is not None:
                out.append((start, c - 1))
                start = None
        if start is not None:
            out.append((start, c1))
        return out

    if not grid or not grid[0]:
        return []
    result = []
    for r0, r1 in split_rows(0, len(grid) - 1, 0, len(grid[0]) - 1):
        for c0, c1 in split_cols(r0, r1, 0, len(grid[0]) - 1):
            for rr0, rr1 in split_rows(r0, r1, c0, c1):
                result.append((rr0, c0, rr1, c1))
    return result


# ---------------------------------------------------------------- table building

def _is_header_row(row: list) -> bool:
    filled = [v for v in row if not _empty(v)]
    if len(filled) < max(2, 0.5 * len(row)):
        return False
    texty = sum(1 for v in filled if _is_text(v) or _looks_period(v))
    # >= 2 distinct labels: a title merged across the row is one repeated value.
    return texty / len(filled) >= 0.8 and len({_header_text(v) for v in filled}) >= 2


def _has_spans(row: list) -> bool:
    """Header rows that group columns ('Q1','Q1','Q2','Q2' after merge-fill, or gaps)."""
    vals = [None if _empty(v) else _header_text(v) for v in row]
    adjacent_dupes = any(a is not None and a == b for a, b in zip(vals, vals[1:]))
    gaps = any(v is None for v in vals[1:])
    return adjacent_dupes or gaps


def _datalike(row: list) -> bool:
    filled = [v for v in row if not _empty(v)]
    return bool(filled) and any(_is_number(v) or isinstance(v, (dt.date, dt.datetime)) for v in filled)


def _dedupe(names: list[str]) -> list[str]:
    used, out = set(), []
    for n in names:
        cand, k = n, 2
        while cand.lower() in used:
            cand, k = f"{n} {k}", k + 1
        used.add(cand.lower())
        out.append(cand)
    return out


def _coerce_column(values: list) -> list:
    """Make a column Parquet-safe: one Python type, or text for mixed columns."""
    kinds = set()
    for v in values:
        if v is None:
            continue
        if isinstance(v, bool):
            kinds.add("bool")
        elif _is_number(v):
            kinds.add("num")
        elif isinstance(v, (dt.datetime, dt.date)):
            kinds.add("date")
        else:
            kinds.add("text")
    if len(kinds) <= 1 and kinds != {"text"}:
        if kinds == {"date"}:
            return [pd.Timestamp(v) if v is not None else None for v in values]
        return values
    out = []
    for v in values:
        if v is None:
            out.append(None)
        elif isinstance(v, float) and v.is_integer():
            out.append(str(int(v)))
        elif isinstance(v, (dt.datetime, dt.date)):
            out.append(v.isoformat())
        else:
            out.append(str(v))
    return out


def _cell_ref(r: int, c: int) -> str:
    return f"{get_column_letter(c + 1)}{r + 1}"


def _build(sheet: _Sheet, box: tuple[int, int, int, int], known_header: bool = False,
           title: str | None = None) -> ParsedTable | None:
    r0, c0, r1, c1 = box
    rows = [sheet.grid[r][c0:c1 + 1] for r in range(r0, r1 + 1)]
    width = c1 - c0 + 1
    stats: dict = {}
    report = {"sheet": sheet.name, "range": f"{_cell_ref(r0, c0)}:{_cell_ref(r1, c1)}",
              "hidden_sheet": sheet.hidden, "notes": []}
    if len(rows) < 2 or width < 2:
        return None

    # 1. Title rows: sparse text rows above the header.
    start = 0
    if not known_header:
        while start < min(len(rows) - 2, 6) and not _is_header_row(rows[start]):
            filled = [v for v in rows[start] if not _empty(v)]
            distinct = list(dict.fromkeys(_header_text(v) for v in filled))
            if len(distinct) > 2 or not all(_is_text(v) for v in filled):
                break
            title = title or " ".join(distinct)
            start += 1
        if start:
            report["notes"].append(f"skipped {start} title row(s)")

    # 2. Header row(s).
    header_rows = 0
    if known_header or _is_header_row(rows[start]):
        header_rows = 1
        if (not known_header and start + 2 < len(rows) and _has_spans(rows[start])
                and _is_header_row(rows[start + 1]) and not _has_spans(rows[start + 1])
                and _datalike(rows[start + 2])):
            header_rows = 2
    if header_rows == 0:
        names = [f"Column {i + 1}" for i in range(width)]
        report["notes"].append("no header row found; columns were numbered")
        raw_headers = [None] * width
    elif header_rows == 1:
        raw_headers = rows[start]
        names = [_header_text(v) if not _empty(v) else f"Column {i + 1}" for i, v in enumerate(raw_headers)]
    else:
        top, sub = rows[start], rows[start + 1]
        filled_top, last = [], None
        for v in top:  # forward-fill group labels across their span
            last = v if not _empty(v) else last
            filled_top.append(last)
        names = []
        for i, (a, b) in enumerate(zip(filled_top, sub)):
            parts = [_header_text(p) for p in (a, b) if not _empty(p)]
            parts = list(dict.fromkeys(parts))
            names.append(" - ".join(parts) if parts else f"Column {i + 1}")
        raw_headers = sub
        report["notes"].append("combined a two-row header")
    names = _dedupe(names)
    report["header_row"] = r0 + start + 1 if header_rows else None
    data = [list(r) for r in rows[start + header_rows:]]

    # 3. Clean data rows.
    header_key = [n.lower() for n in names]
    kept, totals, repeats = [], 0, 0
    for row in data:
        row = [_clean_cell(v, stats) for v in row]
        if all(v is None for v in row):
            continue
        lead = [v for v in row[:3] if isinstance(v, str)]
        if any(_TOTAL_RE.match(v) for v in lead):
            totals += 1
            continue
        if [str(v).lower() if v is not None else "" for v in row] == header_key:
            repeats += 1
            continue
        kept.append(row)
    if totals:
        report["notes"].append(f"removed {totals} total/subtotal row(s)")
    if repeats:
        report["notes"].append(f"removed {repeats} repeated header row(s)")
    if stats.get("error_cells"):
        report["notes"].append(f"{stats['error_cells']} Excel error cell(s) (#N/A, #DIV/0!...) treated as empty")
    if not kept:
        return None

    cols = [[row[i] for row in kept] for i in range(width)]

    # 4. Drop fully empty columns.
    keep_idx = [i for i in range(width) if any(v is not None for v in cols[i])]
    dropped = [names[i] for i in range(width) if i not in keep_idx]
    if dropped:
        report["notes"].append(f"dropped {len(dropped)} empty column(s)")
    names = [names[i] for i in keep_idx]
    raw_headers = [raw_headers[i] for i in keep_idx]
    cols = [cols[i] for i in keep_idx]

    # 5. Report-style grouping: text in leading columns only on the first row of each group.
    filled_down = []
    n = len(kept)
    for i in range(min(3, len(cols))):
        col = cols[i]
        nonnull = [v for v in col if v is not None]
        if n < 4 or not nonnull or col[0] is None or not all(isinstance(v, str) for v in nonnull):
            continue
        blank_frac = 1 - len(nonnull) / n
        # Rows missing the group label must otherwise be real data rows (mostly filled).
        others = [j for j in range(len(cols)) if j != i]
        blank_rows = [r for r in range(n) if col[r] is None]
        rows_filled = sum(1 for r in blank_rows
                          if sum(cols[j][r] is not None for j in others) >= 0.5 * len(others))
        if 0.3 <= blank_frac < 0.95 and others and rows_filled >= 0.8 * len(blank_rows):
            last = None
            filled = []
            for v in col:
                last = v if v is not None else last
                filled.append(last)
            cols[i] = filled
            filled_down.append(names[i])
    if filled_down:
        report["notes"].append(f"filled down grouped values in: {', '.join(filled_down)}")

    frame = pd.DataFrame({name: _coerce_column(col) for name, col in zip(names, cols)})

    # 6. Periods laid out as columns -> unpivot to rows so trends work.
    period_idx = [i for i, h in enumerate(raw_headers) if _looks_period(h)]
    numeric_period = [i for i in period_idx if all(v is None or _is_number(v) for v in cols[i])]
    if len(numeric_period) >= 3 and len(numeric_period) >= 0.5 * len(period_idx):
        period_names = [names[i] for i in numeric_period]
        total_cols = [nm for nm in names if _TOTAL_RE.match(nm)]
        id_cols = [nm for nm in names if nm not in period_names and nm not in total_cols]
        value_name = title.strip() if title and len(title.strip()) <= 60 else "Value"
        if value_name in id_cols or value_name == "Period":
            value_name = "Value"
        long = frame.drop(columns=total_cols).melt(
            id_vars=id_cols, value_vars=period_names, var_name="Period", value_name=value_name)
        long = long[long[value_name].notna()]
        report["notes"].append(
            f"unpivoted {len(period_names)} period columns ({period_names[0]} … {period_names[-1]}) into rows"
            + (f"; ignored total column(s) {', '.join(total_cols)}" if total_cols else ""))
        frame = long.reset_index(drop=True)

    report["rows"], report["columns"] = len(frame), len(frame.columns)
    name = title if title and len(title) <= 60 else sheet.name
    return ParsedTable(name=name, frame=frame, report=report)


def _title_block(grid, box) -> str | None:
    """A small block of label rows (e.g. a merged report title) rather than a table."""
    r0, c0, r1, c1 = box
    if r1 - r0 > 2:
        return None
    first = None
    for r in range(r0, r1 + 1):
        vals = [grid[r][c] for c in range(c0, c1 + 1) if not _empty(grid[r][c])]
        distinct = list(dict.fromkeys(_header_text(v) for v in vals))
        # Titles are one label per row (merged across the width); allow two on a single-row block.
        if len(distinct) > (2 if r0 == r1 else 1) or not all(_is_text(v) for v in vals):
            return None
        first = first or " ".join(distinct)
    return first


# ---------------------------------------------------------------- entry point

def parse_workbook(path: str | Path) -> list[ParsedTable]:
    path = Path(path)
    ext = path.suffix.lower()
    sheets = _load_openpyxl(path) if ext in {".xlsx", ".xlsm", ".xltx", ".xltm"} else _load_pandas(path)

    tables: list[ParsedTable] = []
    for sheet in sheets:
        found: list[ParsedTable] = []
        if sheet.tables:  # Defined Excel Tables are exact; trust them.
            for tname, box in sheet.tables:
                t = _build(sheet, box, known_header=True, title=tname)
                if t:
                    t.report["notes"].insert(0, f"read Excel Table '{tname}'")
                    found.append(t)
        else:
            titles: list[tuple[int, int, int, str]] = []  # (last row, first col, last col, text)
            for box in _blocks(sheet.grid):
                r0, c0, r1, c1 = box
                label = _title_block(sheet.grid, box)
                if label is not None:
                    titles.append((r1, c0, c1, label))
                    continue
                if r1 - r0 < 1 or c1 - c0 < 1:
                    continue
                # A label just above this table (overlapping its columns) names it.
                near = [t for t in titles if r0 - 4 <= t[0] < r0 and t[1] <= c1 and t[2] >= c0]
                title = None
                if near:
                    best = max(near, key=lambda t: t[0])
                    titles.remove(best)
                    title = best[3]
                t = _build(sheet, box, title=title)
                if t:
                    found.append(t)
        if len(found) > 1:
            for i, t in enumerate(found):
                if t.name == sheet.name:
                    t.name = f"{sheet.name} ({i + 1})"
        tables.extend(found)

    return _stack_identical(tables)


def _stack_identical(tables: list[ParsedTable]) -> list[ParsedTable]:
    """Sheets with the same columns (Jan, Feb, Mar… or North, South…) become one table."""
    groups: dict[tuple, list[ParsedTable]] = {}
    for t in tables:
        groups.setdefault(tuple(c.lower() for c in t.frame.columns), []).append(t)
    out = []
    for key, group in groups.items():
        if len(group) < 2 or len(key) < 2:
            out.extend(group)
            continue
        label = "Sheet" if "sheet" not in key else "Source sheet"
        frames = [t.frame.assign(**{label: t.report["sheet"] if t.name == t.report["sheet"] else t.name})
                  for t in group]
        stacked = pd.concat(frames, ignore_index=True)
        names = [t.name for t in group]
        report = {
            "sheet": ", ".join(t.report["sheet"] for t in group),
            "range": "multiple",
            "hidden_sheet": False,
            "notes": [f"stacked {len(group)} tables with identical columns ({', '.join(names[:6])}"
                      f"{'…' if len(names) > 6 else ''}) into one, with a '{label}' column"],
            "parts": [t.report for t in group],
            "rows": len(stacked), "columns": len(stacked.columns),
        }
        out.append(ParsedTable(name=f"{group[0].name} + {len(group) - 1} more", frame=stacked, report=report))
    return out
