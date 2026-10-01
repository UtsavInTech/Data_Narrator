"""Blank Excel templates for a table: the same columns, with dropdowns and checks.

A filled template uploaded back lands in the same table (same columns -> append).
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.worksheet.datavalidation import DataValidation

ROWS = 1000


def build(name: str, fields: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = (name[:28] or "Data").replace("/", "-")
    lists = wb.create_sheet("Lists")
    lists.sheet_state = "hidden"
    guide = wb.create_sheet("How to fill")
    guide.append(["Column", "Required", "What to enter"])
    for c in guide[1]:
        c.font = Font(bold=True)

    list_col = 0
    for i, f in enumerate(fields, 1):
        letter = get_column_letter(i)
        cell = ws.cell(row=1, column=i, value=f["name"])
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="27272A")
        ws.column_dimensions[letter].width = max(len(f["name"]) + 4, 14)
        rng = f"{letter}2:{letter}{ROWS + 1}"
        hint = "Text"
        dv = None
        if f["kind"] == "select" and f["options"]:
            list_col += 1
            lc = get_column_letter(list_col)
            for r, opt in enumerate(f["options"], 1):
                lists.cell(row=r, column=list_col, value=opt)
            dv = DataValidation(type="list", formula1=f"{quote_sheetname('Lists')}!${lc}$1:${lc}${len(f['options'])}",
                                allow_blank=not f["required"], showErrorMessage=False)
            hint = "Pick from the list (a new value is allowed)"
        elif f["kind"] == "number":
            dv = DataValidation(type="whole" if f["integer"] else "decimal", operator="greaterThanOrEqual",
                                formula1="0" if f["min"] == 0 else "-1E+15", allow_blank=not f["required"],
                                showErrorMessage=True, error="Enter a number" + (" (not negative)" if f["min"] == 0 else ""))
            hint = ("Whole number" if f["integer"] else "Number") + (", not negative" if f["min"] == 0 else "")
            for r in range(2, ROWS + 2):
                ws.cell(row=r, column=i).number_format = "0" if f["integer"] else "#,##0.00"
        elif f["kind"] in ("date", "datetime"):
            dv = DataValidation(type="date", operator="greaterThan", formula1="1", allow_blank=not f["required"],
                                showErrorMessage=True, error="Enter a date")
            hint = "Date (DD-MM-YYYY)"
            for r in range(2, ROWS + 2):
                ws.cell(row=r, column=i).number_format = "DD-MM-YYYY"
        elif f["kind"] == "boolean":
            dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=not f["required"])
            hint = "Yes or No"
        if f["unique"]:
            hint += "; must be unique"
        if dv:
            ws.add_data_validation(dv)
            dv.add(rng)
        guide.append([f["name"], "Yes" if f["required"] else "No", hint])
    ws.freeze_panes = "A2"
    guide.column_dimensions["A"].width = 28
    guide.column_dimensions["C"].width = 50
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
