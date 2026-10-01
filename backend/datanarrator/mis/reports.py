"""Daily MIS report: computed once, shown on screen and saved as Excel.

Contents: the day against the previous day with data and the same weekday last
week, month to date, highlights in words, a table per main category, the day's
rows, and checks (problems in the day's data, groups that sent nothing).
A table with a date column is cut by that date; one without uses the batches
added that day.
"""
import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .. import catalog, config, service
from ..db import connect, jsonable, lit, q
from ..engine.fmt import money_or_num, noun_for, pct
from ..engine.insights import pick_dims, pick_measures, pick_time

REPORTS_DIR = config.STORAGE_DIR / "reports"
MAX_ROWS = 50_000
SCREEN_ROWS = 100


def _change(a, b):
    return None if not b else (a - b) / abs(b)


def compute(ds: dict, day: dt.date | None, generated_by: str) -> dict:
    prof = catalog.load_json(ds["id"], "profile")
    if not prof:
        raise ValueError("This table is still being analysed")
    con = connect()
    try:
        rel = service.relation(ds)
        tcol = pick_time(prof)
        measures = [m for m in pick_measures(prof) if m["agg"] in ("sum", "avg")]
        dims = pick_dims(prof)[:3]
        noun = ds["meta"].get("fact_name", ds["name"]) if ds["kind"] == "model" else ds["name"]
        noun = noun_for(noun)

        if tcol and tcol["subtype"] != "year":
            dexpr, src, basis = f"CAST({q(tcol['name'])} AS DATE)", rel, f"by {tcol['name']}"
        else:  # no date column: rows from the batches added each day
            dexpr = "CAST(to_timestamp(__added) AS DATE)"
            parts = {b["part"]: b["created_at"] for b in catalog.batches_for(ds["id"])}
            cases = " ".join(f"WHEN filename LIKE {lit('%' + p)} THEN {t}" for p, t in parts.items()) or "WHEN true THEN 0"
            src = (f"(SELECT * EXCLUDE (filename), CASE {cases} END AS __added FROM read_parquet("
                   f"{lit(catalog.data_glob(ds['id']))}, union_by_name=true, filename=true))")
            basis = "by the day rows were added"

        if day is None:
            day = con.execute(f"SELECT max({dexpr}) FROM {src}").fetchone()[0] or dt.date.today()
        prev = con.execute(f"SELECT max({dexpr}) FROM {src} WHERE {dexpr} < ?", [day]).fetchone()[0]
        week_ago = day - dt.timedelta(days=7)
        month_start = day.replace(day=1)

        metric_sql = ["count(*)"] + [f"{'avg' if m['agg'] == 'avg' else 'sum'}({q(m['name'])})" for m in measures]
        names = [noun] + [m["name"] for m in measures]
        kinds = ["count"] + [m["agg"] for m in measures]

        def window(where, args):
            return con.execute(f"SELECT {', '.join(metric_sql)} FROM {src} WHERE {where}", args).fetchone()

        today_v = window(f"{dexpr} = ?", [day])
        prev_v = window(f"{dexpr} = ?", [prev]) if prev else [None] * len(names)
        week_v = window(f"{dexpr} = ?", [week_ago])
        mtd_v = window(f"{dexpr} BETWEEN ? AND ?", [month_start, day])
        f = lambda v: float(v) if v is not None else None  # noqa: E731
        metrics = [{"label": n, "kind": k, "today": f(t), "previous": f(p), "change": _change(f(t) or 0, f(p)),
                    "week_ago": f(w), "month_to_date": f(m) if k != "avg" else None}
                   for n, k, t, p, w, m in zip(names, kinds, today_v, prev_v, week_v, mtd_v)]

        notes = []
        if day == dt.date.today():
            notes.append("Today is still in progress: figures will grow as entries come in.")
        missing = []
        if prev and dims:
            dn = q(dims[0]["name"])
            missing = [m[0] for m in con.execute(
                f"SELECT DISTINCT CAST({dn} AS VARCHAR) FROM {src} WHERE {dexpr} = ? AND {dn} IS NOT NULL"
                f" EXCEPT SELECT DISTINCT CAST({dn} AS VARCHAR) FROM {src} WHERE {dexpr} = ? ORDER BY 1", [prev, day]).fetchall()]
            if missing:
                notes.append(f"No {noun.lower()} yet from {len(missing)} {dims[0]['name'].lower()}: {', '.join(missing[:8])}.")
        if (today_v[0] or 0) == 0:
            notes.append(f"No {noun.lower()} on {day:%d %b %Y}.")
        elif prev:
            for mt in metrics:
                ch = mt["change"]
                if ch is not None and abs(ch) >= 0.1 and mt["kind"] != "avg":
                    v = money_or_num(mt["today"], mt["label"])
                    notes.append(f"{mt['label']}: {v}, {pct(abs(ch))} {'more' if ch > 0 else 'less'} than {prev:%d %b}.")
                    if len(notes) >= 4:
                        break

        groups = []
        main = next((m for m in measures if m["agg"] == "sum"), None)
        for d in dims:
            dn = q(d["name"])
            ms = [m for m in measures if m["agg"] == "sum"][:2]
            cols = ", ".join(f"sum({q(m['name'])}) FILTER (WHERE {dexpr} = ?), sum({q(m['name'])}) FILTER (WHERE {dexpr} = ?)" for m in ms)
            args = []
            for _ in ms:
                args += [day, prev or day]
            res = con.execute(
                f"SELECT CAST({dn} AS VARCHAR) g, count(*) FILTER (WHERE {dexpr} = ?) n{', ' + cols if cols else ''}"
                f" FROM {src} WHERE {dexpr} IN (?, ?) GROUP BY 1 ORDER BY {3 if ms else 2} DESC NULLS LAST",
                [day, *args, day, prev or day]).fetchall()
            if not res:
                continue
            headers = [d["name"], noun]
            for m in ms:
                headers += [f"{m['name']}", f"{m['name']} ({prev:%d %b})" if prev else f"{m['name']} (prev)"]
            rows = [[r[0] or "(blank)", int(r[1] or 0), *[float(x or 0) for x in r[2:]]] for r in res]
            total = ["Total", sum(r[1] for r in rows), *[sum(r[i] for r in rows) for i in range(2, len(headers))]]
            groups.append({"dimension": d["name"], "headers": headers, "rows": rows, "total": total,
                           "money": [False, False, *[money_or_num(1, m["name"]).startswith("₹") for m in ms for _ in (0, 1)]]})
            if main and len(rows) > 1 and ms:
                leader = max(rows, key=lambda r: r[2])
                if leader[2] > 0:
                    notes.append(f"Top {d['name'].lower()} today: {leader[0]} ({money_or_num(leader[2], ms[0]['name'])}).")

        cur = con.execute(f"SELECT * EXCLUDE (__added) FROM {src} WHERE {dexpr} = ? LIMIT {MAX_ROWS}" if src != rel
                          else f"SELECT * FROM {src} WHERE {dexpr} = ? LIMIT {MAX_ROWS}", [day])
        row_cols = [c[0] for c in cur.description]
        row_data = [list(r) for r in cur.fetchall()]

        checks = []
        if missing:
            checks.append({"column": dims[0]["name"], "problem": "Sent nothing this day: " + ", ".join(missing[:20]), "rows": len(missing)})
        for c in prof["columns"]:
            cn = q(c["name"])
            if c["role"] == "measure" and c["min"] is not None and float(c["min"]) >= 0:
                n = con.execute(f"SELECT count(*) FROM {src} WHERE {dexpr} = ? AND {cn} < 0", [day]).fetchone()[0]
                if n:
                    checks.append({"column": c["name"], "problem": "Negative values", "rows": n})
            if c["null_pct"] < 1 and c["role"] not in ("text", "empty"):
                n = con.execute(f"SELECT count(*) FROM {src} WHERE {dexpr} = ? AND {cn} IS NULL", [day]).fetchone()[0]
                if n:
                    checks.append({"column": c["name"], "problem": "Missing values", "rows": n})
            if c["role"] == "dimension" and prev and c["subtype"] != "status":
                new = con.execute(
                    f"SELECT DISTINCT CAST({cn} AS VARCHAR) FROM {src} WHERE {dexpr} = ? AND {cn} IS NOT NULL"
                    f" AND {cn} NOT IN (SELECT {cn} FROM {src} WHERE {dexpr} < ? AND {cn} IS NOT NULL) LIMIT 10",
                    [day, day]).fetchall()
                if new:
                    checks.append({"column": c["name"], "problem": "New values: " + ", ".join(v[0] for v in new), "rows": len(new)})
            if c["role"] == "id" and c["unique"] > 0.9 * prof["sample_rows"]:
                n = con.execute(f"SELECT count(*) - count(DISTINCT {cn}) FROM {src} WHERE {dexpr} = ?", [day]).fetchone()[0]
                if n:
                    checks.append({"column": c["name"], "problem": "Repeated IDs", "rows": n})
        if checks:
            notes.append(f"{len(checks)} check{'s' if len(checks) > 1 else ''} need attention.")
    finally:
        con.close()

    return {"title": f"{ds['name']}: daily report", "dataset": ds["name"], "day": day.isoformat(),
            "previous_day": prev.isoformat() if prev else None, "week_ago": week_ago.isoformat(), "basis": basis,
            "generated_by": generated_by, "generated_at": dt.datetime.now().isoformat(timespec="minutes"),
            "metrics": metrics, "highlights": notes or ["Nothing unusual."], "groups": groups,
            "rows": {"columns": row_cols, "data": row_data}, "checks": checks}


# ---------------------------------------------------------------- Excel

_HEAD = PatternFill("solid", fgColor="1C1B19")
_TOTAL = PatternFill("solid", fgColor="F1F0EC")
_LINE = Border(bottom=Side(style="thin", color="E6E3DC"))
_NUM = "#,##0.##"


def _sheet_table(ws, start_row: int, headers: list[str], rows: list[list], total: list | None = None):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=c, value=h)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), _HEAD
        cell.alignment = Alignment(vertical="center")
    for r, row in enumerate(rows, start_row + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = _LINE
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                cell.number_format = _NUM
            elif isinstance(v, dt.datetime):
                cell.number_format = "DD-MM-YYYY" if v.time() == dt.time() else "DD-MM-YYYY HH:MM"
            elif isinstance(v, dt.date):
                cell.number_format = "DD-MM-YYYY"
    if total:
        r = start_row + len(rows) + 1
        for c, v in enumerate(total, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font, cell.fill = Font(bold=True), _TOTAL
            if isinstance(v, (int, float)):
                cell.number_format = _NUM
    for c in range(1, len(headers) + 1):
        width = max([len(str(headers[c - 1]))] + [len(f"{row[c - 1]:,.2f}" if isinstance(row[c - 1], float) else str(row[c - 1] or ""))
                                                   for row in rows[:500]])
        ws.column_dimensions[get_column_letter(c)].width = min(max(width + 2, 10), 45)


def write_excel(rep: dict, ds_id: str) -> Path:
    day = dt.date.fromisoformat(rep["day"])
    prev = dt.date.fromisoformat(rep["previous_day"]) if rep["previous_day"] else None
    week = dt.date.fromisoformat(rep["week_ago"])
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = rep["title"]
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"{day:%A, %d %B %Y}  ·  {rep['basis']}  ·  generated {rep['generated_at'].replace('T', ' ')} by {rep['generated_by']}"
    ws["A2"].font = Font(color="7A776F")
    rows = [[m["label"] if m["kind"] != "avg" else f"{m['label']} (average)", m["today"], m["previous"],
             f"{m['change']:+.1%}" if m["change"] is not None else "", m["week_ago"], m["month_to_date"]] for m in rep["metrics"]]
    _sheet_table(ws, 4, ["Metric", f"{day:%d %b %Y}", f"Previous ({prev:%d %b})" if prev else "Previous", "Change",
                         f"Same day last week ({week:%d %b})", "Month to date"], rows)
    r = 6 + len(rows)
    ws.cell(row=r, column=1, value="Highlights").font = Font(bold=True, size=12)
    for i, n in enumerate(rep["highlights"], 1):
        ws.cell(row=r + i, column=1, value=f"•  {n}")
    ws.column_dimensions["A"].width = 34
    for g in rep["groups"]:
        sh = wb.create_sheet(f"By {g['dimension']}"[:31])
        _sheet_table(sh, 1, g["headers"], g["rows"], g["total"])
        sh.freeze_panes = "B2"
    sh = wb.create_sheet("Rows")
    _sheet_table(sh, 1, rep["rows"]["columns"], rep["rows"]["data"])
    sh.freeze_panes = "A2"
    if rep["rows"]["data"]:
        sh.auto_filter.ref = f"A1:{get_column_letter(len(rep['rows']['columns']))}{len(rep['rows']['data']) + 1}"
    sh = wb.create_sheet("Checks")
    checks = [[c["column"], c["problem"], c["rows"]] for c in rep["checks"]]
    _sheet_table(sh, 1, ["Column", "Problem", "Rows"], checks or [["-", "No problems found", 0]])
    out_dir = REPORTS_DIR / ds_id
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{day:%Y-%m-%d}.xlsx"
    wb.save(path)
    return path


def build(ds: dict, day: dt.date | None, generated_by: str) -> tuple[Path, dt.date]:
    rep = compute(ds, day, generated_by)
    return write_excel(rep, ds["id"]), dt.date.fromisoformat(rep["day"])


def for_screen(rep: dict) -> dict:
    """JSON-safe, with the day's rows trimmed for display."""
    out = dict(rep)
    out["rows"] = {"columns": rep["rows"]["columns"], "total": len(rep["rows"]["data"]),
                   "data": [[jsonable(v) for v in r] for r in rep["rows"]["data"][:SCREEN_ROWS]]}
    return out


def list_reports(owner: str) -> list[dict]:
    names = {d["id"]: d["name"] for d in catalog.list_for(owner)}
    out = []
    for p in REPORTS_DIR.glob("*/*.xlsx"):
        if p.parent.name in names:
            out.append({"ds_id": p.parent.name, "dataset": names[p.parent.name], "date": p.stem,
                        "file": f"{p.parent.name}/{p.name}", "created_at": p.stat().st_mtime})
    return sorted(out, key=lambda r: (r["date"], r["created_at"]), reverse=True)
